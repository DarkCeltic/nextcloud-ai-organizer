#!/usr/bin/env python3

import argparse
import logging
import sys
from pathlib import Path
from typing import Any, Dict, Iterable, Optional
from openpyxl import load_workbook

from python_organizer_local_llm.classifier import Classifier
from python_organizer_local_llm.database import Database
from python_organizer_local_llm.nextcloud import NextcloudClient
from python_organizer_local_llm.scanner import Scanner
from python_organizer_local_llm.sensitive import sensitive_filename, safe_suggestion
from python_organizer_local_llm.settings import EnvironmentSettings, load_environment_settings

DEFAULT_CONFIG = "config.yaml"


def configure_logging(verbose: bool = False) -> None:
    """Configure console logging."""
    level = logging.DEBUG if verbose else logging.INFO

    logging.basicConfig(
        level=level,
        format="%(asctime)s | %(levelname)-8s | %(name)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


class Organizer:
    """
    Main coordinator for the Nextcloud AI Organizer.

    Workflow:
        Nextcloud
            -> Scanner
            -> Extract readable text
            -> AI classifier
            -> SQLite suggestions database

    This version is intentionally suggestion-only.

    It does NOT:
      - Rename files
      - Move files
      - Delete files
      - Apply tags
      - Send files to Paperless
    """

    def __init__(
        self,
        config_file: str = DEFAULT_CONFIG,
        force: bool = False,
        limit: Optional[int] = None,
        runtime_settings: Optional[EnvironmentSettings] = None,
    ) -> None:
        self.log = logging.getLogger(self.__class__.__name__)
        self.config_file = config_file
        self.force = force
        self.limit = limit
        # FastAPI/CLI entry points load .env once and pass the resulting runtime
        # settings here. Direct library users may rely on real environment variables.
        self.runtime_settings = runtime_settings or load_environment_settings(
            load_env_file=False
        )

        if self.limit is not None and self.limit < 1:
            raise ValueError("--limit must be greater than zero.")

        config_path = Path(config_file)
        if not config_path.is_file():
            raise FileNotFoundError(
                f"Configuration file does not exist: {config_file}"
            )

        # These constructors must match the companion modules.
        # If your existing modules use different constructor signatures,
        # upload them with this file and they can be matched exactly.
        self.database = Database(config_file)
        self.nextcloud = NextcloudClient(
            config_file, runtime_settings=self.runtime_settings
        )
        self.classifier = Classifier(
            config_file, runtime_settings=self.runtime_settings
        )

        paperless_config = self.classifier.config.get("paperless", {})
        self.paperless_enabled = self.classifier.paperless_enabled
        self.paperless_inbox = str(
            paperless_config.get("inbox_path", "/inbox")
        ).strip() or "/inbox"
        if not self.paperless_inbox.startswith("/"):
            self.paperless_inbox = "/" + self.paperless_inbox

        self.scanner = Scanner(
            nextcloud=self.nextcloud,
            database=self.database,
            config_file=config_file,
        )

    def run(self) -> int:
        """Scan eligible files and generate pending AI suggestions."""
        self.database.initialize()
        # CLI runs use the same persisted SQLite selections as the ExApp UI.
        from python_organizer_local_llm.settings import SettingsService
        SettingsService(self)

        processed = 0
        skipped = 0
        failed = 0

        for file_info in self._scan_files():
            try:
                result = self._process_file(file_info)

                if result == "processed":
                    processed += 1
                else:
                    skipped += 1

            except Exception:
                failed += 1
                path = self._get_value(file_info, "path", default="(unknown)")
                self.log.exception("Failed to process file: %s", path)

        self._print_summary(
            processed=processed,
            skipped=skipped,
            failed=failed,
        )

        return 1 if failed else 0

    def _scan_files(self) -> Iterable[Dict[str, Any]]:
        """
        Return eligible files from Scanner.

        `scan()` is the preferred interface. `scan_files()` is accepted as a
        compatibility fallback because earlier versions of this project used
        both names while the scanner module was being developed.
        """
        if hasattr(self.scanner, "scan"):
            return self.scanner.scan(
                force=self.force,
                limit=self.limit,
            )

        if hasattr(self.scanner, "scan_files"):
            return self.scanner.scan_files(
                force=self.force,
                limit=self.limit,
            )

        raise AttributeError(
            "Scanner must provide scan(force=..., limit=...) "
            "or scan_files(force=..., limit=...)."
        )

    @staticmethod
    def _get_value(
        file_info: Any,
        *names: str,
        default: Any = None,
    ) -> Any:
        """Read a field from either a dictionary or an object."""
        if isinstance(file_info, dict):
            for name in names:
                if name in file_info:
                    return file_info[name]
            return default

        for name in names:
            if hasattr(file_info, name):
                return getattr(file_info, name)

        return default

    def _process_file(self, file_info: Any) -> str:
        """Generate and store one suggestion for a scanned file."""
        path = self._get_value(file_info, "path")
        file_id = self._get_value(
            file_info,
            "file_id",
            "nextcloud_file_id",
            "id",
        )
        etag = self._get_value(file_info, "etag", default="")
        mime_type = self._get_value(
            file_info,
            "mime_type",
            "mimetype",
            "content_type",
            default="",
        )
        size = self._get_value(file_info, "size", default=0)

        if not path:
            raise ValueError("Scanner result is missing a file path.")

        if file_id is None:
            raise ValueError(
                f"Scanner result is missing a Nextcloud file ID: {path}"
            )

        # Enforce the current selection even when called outside Scanner.scan().
        if not self.scanner.is_supported_file(path):
            return "skipped"

        database_file_id = self.database.upsert_file(
            nextcloud_file_id=file_id,
            path=path,
            etag=etag,
            mime_type=mime_type,
            size=size,
        )

        if database_file_id is None:
            raise RuntimeError(
                "Database.upsert_file() did not return a database file ID "
                f"for {path}"
            )

        # Extensionless credentials and secret-like names receive safe metadata,
        # never file-content downloads or local-LLM requests (including CLI runs).
        if sensitive_filename(path):
            suggestion = safe_suggestion(path)
            self._validate_suggestion(suggestion)
            suggestion_id = self.database.save_suggestion(
                file_id=database_file_id,
                suggested_filename=suggestion['suggested_filename'],
                suggested_folder=suggestion['suggested_folder'],
                tags=suggestion['tags'],
                category=suggestion['category'],
                paperless_candidate=False,
                confidence=0.0,
                reason=suggestion['reason'],
                manual_review_only=True,
            )
            self.database.mark_processed(database_file_id)
            self._print_suggestion(path, suggestion_id, suggestion)
            return 'processed'

        # Retrieve and extract readable content.
        try:
            extractor = getattr(self.nextcloud, 'get_file_text_with_metadata', None)
            extraction = (extractor(path, file_id=str(file_id), etag=str(etag or ''))
                          if callable(extractor) else
                          {'text': self.nextcloud.get_file_text(path), 'ocr_used': False})
            content = extraction['text']
            if not content or not content.strip():
                raise ValueError('No readable text was extracted; review the file manually.')
        except Exception as exc:
            # A failed extraction remains available in Failed and must not be
            # silently marked processed or retried every scheduler interval.
            self.database.record_analysis_failure(
                str(file_id), path, str(exc), 'ocr' if path.lower().endswith('.pdf') else 'extract'
            )
            self.log.warning('Extraction failed; file recorded in Failed: %s', path)
            return "skipped"

        # Existing Nextcloud structure helps prevent the model from inventing
        # unnecessary folders.
        try:
            existing_folders = self.nextcloud.get_folder_tree()
        except Exception:
            self.log.exception("Unable to retrieve folder tree.")
            existing_folders = []

        # Existing tags are optional.
        try:
            existing_tags = self.nextcloud.get_tags()
        except Exception:
            self.log.exception("Unable to retrieve Nextcloud tags.")
            existing_tags = []

        filename = Path(path).name

        # Ask the local LLM for suggestions.
        suggestion = self.classifier.classify(
            filename=filename,
            file_path=path,
            content=content,
            existing_folders=existing_folders,
            existing_tags=existing_tags,
            force_nextcloud=False,
        )

        if not suggestion:
            raise RuntimeError(
                f"Classifier returned no suggestion for {path}"
            )

        # Defense in depth: apply the classifier's deterministic Paperless
        # routing policy again before storing anything. This protects against
        # future code paths that might bypass Classifier.classify() normalization.
        suggestion = self.classifier._apply_paperless_policy(
            suggestion=suggestion,
            filename=filename,
            content=content,
        )

        self._validate_suggestion(suggestion)

        # Store suggestion only.
        suggestion_id = self.database.save_suggestion(
            file_id=database_file_id,
            suggested_filename=suggestion.get("suggested_filename"),
            suggested_folder=suggestion.get("suggested_folder"),
            tags=suggestion.get("tags", []),
            category=suggestion.get("category", "unknown"),
            paperless_candidate=suggestion.get(
                "paperless_candidate",
                False,
            ),
            confidence=suggestion.get("confidence", 0.0),
            reason=suggestion.get("reason", ""),
            status="pending",
            manual_review_only=bool(extraction["ocr_used"]),
            ocr_used=bool(extraction["ocr_used"]),
        )

        self.database.mark_processed(database_file_id)
        self.database.resolve_analysis_failure(str(file_id))

        self._print_suggestion(
            original_path=path,
            suggestion_id=suggestion_id,
            suggestion=suggestion,
        )

        return "processed"

    def _validate_suggestion(self, suggestion: dict) -> None:
        """Perform basic safety checks before storing AI output."""
        if not isinstance(suggestion, dict):
            raise ValueError(
                "Classifier response must be a dictionary."
            )

        required_fields = [
            "suggested_filename",
            "suggested_folder",
            "tags",
            "category",
            "paperless_candidate",
            "confidence",
            "reason",
        ]

        for field in required_fields:
            if field not in suggestion:
                raise ValueError(
                    "Classifier response is missing required field: "
                    f"{field}"
                )

        filename = suggestion.get("suggested_filename", "")
        if not isinstance(filename, str) or not filename.strip():
            raise ValueError(
                "Suggested filename must be a non-empty string."
            )

        invalid_filename_characters = ["/", "\\", "\x00"]
        for character in invalid_filename_characters:
            if character in filename:
                raise ValueError(
                    "Suggested filename contains invalid character: "
                    f"{character!r}"
                )

        folder = suggestion.get("suggested_folder", "")
        if not isinstance(folder, str) or not folder.strip():
            raise ValueError(
                "Suggested folder must be a non-empty string."
            )

        if not folder.startswith("/"):
            raise ValueError(
                "Suggested folder must begin with '/'."
            )

        tags = suggestion.get("tags", [])
        if not isinstance(tags, list):
            raise ValueError("Suggested tags must be a list.")

        if not all(isinstance(tag, str) for tag in tags):
            raise ValueError(
                "Every suggested tag must be a string."
            )

        paperless_candidate = suggestion.get("paperless_candidate")
        if not isinstance(paperless_candidate, bool):
            raise ValueError(
                "paperless_candidate must be true or false."
            )

        try:
            confidence = float(suggestion.get("confidence", 0.0))
        except (TypeError, ValueError) as exc:
            raise ValueError(
                "Confidence must be numeric."
            ) from exc

        if confidence < 0.0 or confidence > 1.0:
            raise ValueError(
                "Confidence must be between 0.0 and 1.0."
            )

    def _print_suggestion(
        self,
        original_path: str,
        suggestion_id: int,
        suggestion: dict,
    ) -> None:
        """Print a human-readable review summary."""
        filename = suggestion.get("suggested_filename", "(none)")
        folder = suggestion.get("suggested_folder", "(none)")
        tags = suggestion.get("tags", [])
        category = suggestion.get("category", "unknown")
        confidence = suggestion.get("confidence", 0.0)
        paperless = suggestion.get("paperless_candidate", False)
        reason = suggestion.get("reason", "")

        try:
            confidence_text = f"{float(confidence) * 100:.1f}%"
        except (TypeError, ValueError):
            confidence_text = str(confidence)

        print()
        print("=" * 72)
        print(f"SUGGESTION ID: {suggestion_id}")
        print(f"FILE:          {original_path}")
        print(f"FILENAME:      {filename}")
        if paperless:
            print(f"DESTINATION:   {folder}")
        else:
            print(f"FOLDER:        {folder}")
            print("TAGS:          " + (", ".join(tags) if tags else "(none)"))
            print(f"CATEGORY:      {category}")
        print(f"CONFIDENCE:    {confidence_text}")
        print("PAPERLESS:     " + ("Yes" if paperless else "No"))

        if reason:
            print(f"REASON:        {reason}")

        print("STATUS:        Pending review")
        print("=" * 72)
        print()

    def _print_summary(
        self,
        processed: int,
        skipped: int,
        failed: int,
    ) -> None:
        """Print end-of-run totals."""
        self.log.info("Organizer finished.")
        self.log.info("Processed: %d", processed)
        self.log.info("Skipped:   %d", skipped)
        self.log.info("Failed:    %d", failed)


def build_argument_parser() -> argparse.ArgumentParser:
    """Build command-line options."""
    parser = argparse.ArgumentParser(
        description=(
            "AI-assisted Nextcloud file python_organizer_local_llm. "
            "This version generates suggestions only."
        )
    )

    parser.add_argument(
        "-c",
        "--config",
        default=DEFAULT_CONFIG,
        help=(
            "Path to configuration file "
            f"(default: {DEFAULT_CONFIG})"
        ),
    )

    parser.add_argument(
        "--force",
        action="store_true",
        help=(
            "Reprocess files even if their Nextcloud ETag has not changed."
        ),
    )

    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Maximum number of eligible files to process during this run.",
    )

    parser.add_argument(
        "--scan-summary",
        action="store_true",
        help=(
            "Show scanner statistics without running AI classification."
        ),
    )

    parser.add_argument(
        "--pending",
        action="store_true",
        help="Display pending suggestions from the SQLite database.",
    )

    parser.add_argument(
        "-v",
        "--verbose",
        action="store_true",
        help="Enable debug logging.",
    )

    return parser


def print_scan_summary(organizer: Organizer) -> int:
    """Display scanner counts without using Ollama."""
    organizer.database.initialize()
    from python_organizer_local_llm.settings import SettingsService
    SettingsService(organizer)

    summary = organizer.scanner.get_scan_summary(
        force=organizer.force
    )

    print()
    print("Nextcloud AI Organizer Scan Summary")
    print("=" * 40)
    print(f"Total files:     {summary.get('total', 0)}")
    print(f"Eligible:        {summary.get('eligible', 0)}")
    print(f"Unchanged:       {summary.get('unchanged', 0)}")
    print(f"Excluded:        {summary.get('excluded', 0)}")
    print(f"Unsupported:     {summary.get('unsupported', 0)}")
    print(f"Ignored:         {summary.get('ignored', 0)}")
    print()

    return 0


def print_pending(organizer: Organizer) -> int:
    """Display pending AI suggestions."""
    organizer.database.initialize()

    suggestions = organizer.database.list_pending(limit=100)

    if not suggestions:
        print("No pending suggestions.")
        return 0

    print()
    print(f"Pending suggestions: {len(suggestions)}")
    print("=" * 72)

    for suggestion in suggestions:
        print(f"ID:         {suggestion.get('id')}")
        print(f"File:       {suggestion.get('original_path')}")
        print(f"Filename:   {suggestion.get('suggested_filename')}")

        paperless = bool(suggestion.get("paperless_candidate"))
        if paperless:
            print(f"Destination:{suggestion.get('suggested_folder'):>12}")
        else:
            print(f"Folder:     {suggestion.get('suggested_folder')}")
            tags = suggestion.get("tags", [])
            print("Tags:       " + (", ".join(tags) if tags else "(none)"))
            print(f"Category:   {suggestion.get('category')}")

        confidence = suggestion.get("confidence", 0.0)
        try:
            confidence_text = f"{float(confidence) * 100:.1f}%"
        except (TypeError, ValueError):
            confidence_text = str(confidence)

        print(f"Confidence: {confidence_text}")
        print(
            "Paperless:  "
            + (
                "Yes"
                if suggestion.get("paperless_candidate")
                else "No"
            )
        )
        print(f"Reason:     {suggestion.get('reason', '')}")
        print("-" * 72)

    print()
    return 0

def extract_xlsx(file_path):
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    content = []

    try:
        for sheet in workbook.worksheets:
            content.append(f"Worksheet: {sheet.title}")

            for row in sheet.iter_rows(
                max_row=20,
                values_only=True
            ):
                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]

                if values:
                    content.append(" | ".join(values))

        return "\n".join(content)

    finally:
        workbook.close()


def main() -> int:
    """Application entry point."""
    parser = build_argument_parser()
    args = parser.parse_args()

    configure_logging(args.verbose)
    log = logging.getLogger("python_organizer_local_llm")

    try:
        organizer = Organizer(
            config_file=args.config,
            force=args.force,
            limit=args.limit,
            runtime_settings=load_environment_settings(load_env_file=True),
        )

        if args.scan_summary:
            return print_scan_summary(organizer)

        if args.pending:
            return print_pending(organizer)

        return organizer.run()

    except KeyboardInterrupt:
        log.warning("Organizer interrupted.")
        return 130

    except Exception:
        log.exception("Organizer failed.")
        return 1


if __name__ == "__main__":
    sys.exit(main())

