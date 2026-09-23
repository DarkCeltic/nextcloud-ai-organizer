#!/usr/bin/env python3

import io
import os
import logging
import mimetypes
import re
import xml.etree.ElementTree as ET
from pathlib import PurePosixPath
from typing import Dict, List, Optional
from urllib.parse import quote, urljoin

import requests
import yaml

from python_organizer_local_llm.sensitive import sensitive_filename, sensitive_content
from python_organizer_local_llm.ocr import recognize_pdf, OCRProcessingError
from python_organizer_local_llm.file_types import BY_EXTENSION


class NextcloudClient:
    """
    Direct Nextcloud client for the AI python_organizer_local_llm.

    Uses:
      - WebDAV for listing and downloading files
      - WebDAV for reading and applying file operations
      - System Tags WebDAV API for reading and assigning collaborative tags
    """

    DAV_NS = "DAV:"
    NC_NS = "http://nextcloud.org/ns/"

    def __init__(self, config_file: str = "config.yaml"):
        self.log = logging.getLogger("nextcloud")
        self.config = self._load_config(config_file)

        cfg = self.config.get("nextcloud", {})

        self.base_url = str(os.getenv("NEXTCLOUD_INTERNAL_URL") or cfg.get("url", "")).rstrip("/")
        self.username = str(os.getenv("NEXTCLOUD_USERNAME") or cfg.get("username", ""))
        self.password = str(os.getenv("NEXTCLOUD_APP_PASSWORD") or cfg.get("app_password", cfg.get("password", "")))
        self.verify_ssl = bool(cfg.get("verify_ssl", True))
        self.timeout = int(cfg.get("timeout", 60))

        if not self.base_url:
            raise RuntimeError("nextcloud.url is required in config.yaml")

        if not self.username:
            raise RuntimeError("nextcloud.username is required in config.yaml")

        if not self.password:
            raise RuntimeError(
                "nextcloud.app_password is required in config.yaml. "
                "Use a Nextcloud app password rather than your normal password."
            )

        organizer_cfg = self.config.get("python_organizer_local_llm", {})
        scanner_cfg = self.config.get("scanner", {})

        self.scan_paths = scanner_cfg.get(
            "scan_paths",
            organizer_cfg.get("scan_paths", ["/AI Inbox"]),
        )

        self.exclude_paths = scanner_cfg.get(
            "exclude_paths",
            organizer_cfg.get(
                "exclude_paths",
                ["/paperless-media", "/inbox", "/Photos", "/AI Ignored"],
            ),
        )

        self.allowed_extensions = {
            self._normalize_extension(ext)
            for ext in scanner_cfg.get(
                "allowed_extensions",
                ["pdf", "txt", "md", "rtf", "docx", "odt"],
            )
            if self._normalize_extension(ext) in BY_EXTENSION
        }

        self.ocr_enabled = True
        self.ocr_max_pages = 10
        # Volatile cache only: never write extracted (possibly private) text to SQLite.
        # File ID + ETag prevents reusing an obsolete OCR result after modification.
        self._ocr_cache = {}
        self.session = requests.Session()
        self.session.auth = (self.username, self.password)
        self.session.headers.update(
            {
                "User-Agent": "Nextcloud-AI-Organizer/0.1",
                "OCS-APIRequest": "true",
            }
        )

    def get_scan_paths(self) -> List[str]:
        return [self._normalize_path(path) for path in self.scan_paths]

    def should_exclude(self, path: str) -> bool:
        path = self._normalize_path(path)

        for excluded in self.exclude_paths:
            excluded = self._normalize_path(excluded)

            if path == excluded or path.startswith(excluded.rstrip("/") + "/"):
                return True

        return False

    def is_supported_file(self, path: str) -> bool:
        suffix = PurePosixPath(path).suffix.lower()
        return bool((suffix and suffix in self.allowed_extensions) or
                    (getattr(self, "allow_sensitive", True) and sensitive_filename(path)))

    def list_files_recursive(self, path: str) -> List[Dict]:
        """
        Recursively list files below a Nextcloud path.

        Returns dictionaries compatible with scanner.py and python_organizer_local_llm.py.
        """
        root = self._normalize_path(path)
        files = []

        for item in self._propfind(root, depth="infinity"):
            item_path = item.get("path")

            if not item_path or item_path == root:
                continue

            if item.get("is_directory"):
                continue

            files.append(item)

        return files

    def get_folder_tree(self) -> List[str]:
        """
        Return all folder paths visible below the configured scan roots'
        parent areas.

        To keep this predictable, folder_tree_paths can be set in config.
        If omitted, the root Files directory is scanned for folders.
        """
        cfg = self.config.get("nextcloud", {})
        roots = cfg.get("folder_tree_paths", ["/"])

        folders = []

        for root in roots:
            try:
                for item in self._propfind(
                        self._normalize_path(root),
                        depth="infinity",
                ):
                    if item.get("is_directory"):
                        path = item.get("path")

                        if not path:
                            continue

                        # Do not offer these storage locations to the LLM.
                        blocked_paths = (
                            "/paperless-media",
                            "/inbox",
                            "/AI Inbox",
                            "/AI Ignored",
                        )

                        normalized = path.rstrip("/").lower()

                        if any(
                                normalized == blocked.lower()
                                or normalized.startswith(blocked.lower() + "/")
                                for blocked in blocked_paths
                        ):
                            continue

                        if path not in folders:
                            folders.append(path)
            except Exception:
                self.log.exception("Could not read folder tree below %s", root)

        return sorted(set(folders), key=str.lower)

    def get_tags(self) -> list[str]:
        """Read actual collaborative tag names from the system-tag DAV properties.

        get_system_tags() parses oc:display-name; d:displayname is not a reliable
        name field for these records. No model should be given tag IDs as names.
        """
        try:
            names = {tag['name'] for tag in self.get_system_tags() if tag.get('name')}
            return sorted(names, key=str.casefold)
        except Exception as exc:
            self.log.warning('Unable to retrieve Nextcloud tags: %s', exc)
            return []

    def find_file_by_id(self, file_id: str, root: str = "/") -> Optional[Dict]:
        """Find a file by Nextcloud numeric file ID.

        This is a fallback for ExApp requests whose short-lived file-action
        context is unavailable. It recursively scans the configured user's
        DAV tree and returns the first matching file.
        """
        wanted = str(file_id or "").strip()
        if not wanted:
            return None

        for item in self._propfind(self._normalize_path(root), depth="infinity"):
            if str(item.get("file_id") or "") == wanted:
                return item
        return None

    def ensure_folder(self, path: str) -> str:
        """Create missing folders below the user's Files root."""
        path = self._normalize_path(path)
        if path == "/":
            return path

        current = ""
        for part in path.strip("/").split("/"):
            current += "/" + part
            response = self.session.request(
                "MKCOL",
                self._dav_url(current),
                timeout=self.timeout,
                verify=self.verify_ssl,
            )
            if response.status_code in (201, 405):
                continue
            response.raise_for_status()
        return path

    def move_file(self, source: str, destination: str, overwrite: bool = False) -> str:
        """Move or rename a file using WebDAV MOVE."""
        source = self._normalize_path(source)
        destination = self._normalize_path(destination)
        if source == destination:
            return destination

        response = self.session.request(
            "MOVE",
            self._dav_url(source),
            headers={
                "Destination": self._dav_url(destination),
                "Overwrite": "T" if overwrite else "F",
            },
            timeout=self.timeout,
            verify=self.verify_ssl,
        )
        if response.status_code in (201, 204):
            return destination
        if response.status_code == 412:
            raise FileExistsError(f"Destination already exists: {destination}")
        response.raise_for_status()
        return destination

    def get_system_tags(self) -> List[Dict]:
        """Return visible system tags with their IDs and names."""
        url = f"{self.base_url}/remote.php/dav/systemtags/"
        body = """<?xml version="1.0" encoding="UTF-8"?>
<d:propfind xmlns:d="DAV:" xmlns:oc="http://owncloud.org/ns">
  <d:prop>
    <oc:id />
    <oc:display-name />
    <d:orderby/>
    <oc:user-visible />
    <oc:user-assignable />
    <oc:can-assign />
  </d:prop>
</d:propfind>
"""
        response = self.session.request(
            "PROPFIND",
            url,
            headers={"Depth": "1", "Content-Type": "application/xml"},
            data=body,
            timeout=self.timeout,
            verify=self.verify_ssl,
        )
        response.raise_for_status()
        root = ET.fromstring(response.content)
        dav = {"d": "DAV:", "oc": "http://owncloud.org/ns/"}
        results = []
        for node in root.findall("d:response", dav):
            tag_id = node.findtext(".//oc:id", default="", namespaces=dav).strip()
            name = node.findtext(".//oc:display-name", default="", namespaces=dav).strip()
            if tag_id and name:
                results.append({"id": tag_id, "name": name})
        return results

    def create_system_tag(self, name: str) -> str:
        """Create a normal visible/assignable Nextcloud system tag."""
        name = str(name or "").strip()
        if not name:
            raise ValueError("Tag name cannot be empty.")
        response = self.session.post(
            f"{self.base_url}/remote.php/dav/systemtags/",
            json={
                "name": name,
                "userVisible": True,
                "userAssignable": True,
                "canAssign": True,
            },
            timeout=self.timeout,
            verify=self.verify_ssl,
        )
        if response.status_code not in (200, 201):
            response.raise_for_status()

        location = response.headers.get("Content-Location") or response.headers.get("Location") or ""
        if location:
            return location.rstrip("/").split("/")[-1]

        # Some Nextcloud versions do not return the ID in a convenient header.
        for tag in self.get_system_tags():
            if tag["name"].casefold() == name.casefold():
                return str(tag["id"])
        raise RuntimeError(f"Created tag but could not resolve its ID: {name}")

    def assign_tags(self, file_id: str, tags: List[str], create_missing: bool = True) -> None:
        """Assign collaborative/system tags to a Nextcloud file ID."""
        existing = {tag["name"].casefold(): tag for tag in self.get_system_tags()}
        for raw_name in tags:
            name = str(raw_name or "").strip()
            if not name:
                continue
            tag = existing.get(name.casefold())
            if tag is None:
                if not create_missing:
                    raise KeyError(f"Nextcloud system tag does not exist: {name}")
                tag_id = self.create_system_tag(name)
                tag = {"id": tag_id, "name": name}
                existing[name.casefold()] = tag

            url = (
                f"{self.base_url}/remote.php/dav/systemtags-relations/files/"
                f"{quote(str(file_id), safe='')}/{quote(str(tag['id']), safe='')}"
            )
            response = self.session.put(
                url,
                timeout=self.timeout,
                verify=self.verify_ssl,
            )
            # 201/204 are success; 409 can mean the relationship already exists.
            if response.status_code not in (201, 204, 409):
                response.raise_for_status()

    def get_file_text_with_metadata(self, path: str, *, file_id: str = '',
                                    etag: str = '') -> dict:
        """Download once; OCR only when a PDF lacks a searchable text layer."""
        path = self._normalize_path(path)
        suffix = PurePosixPath(path).suffix.lower()
        cache_key = (str(file_id), str(etag).strip('"')) if file_id and etag else None
        if suffix == '.pdf' and self.ocr_enabled and cache_key in self._ocr_cache:
            return {'text': self._ocr_cache[cache_key], 'ocr_used': True}
        content = self.download_file(path)
        if suffix == '.pdf':
            try:
                text = self._extract_pdf(content)
            except Exception as exc:
                raise OCRProcessingError('OCR: PDF text extraction failed; check whether it is damaged or protected.') from exc
            if text and text.strip():
                return {'text': text, 'ocr_used': False}
            if not self.ocr_enabled:
                raise OCRProcessingError('OCR: PDF has no readable text and OCR is disabled in Settings.')
            text = recognize_pdf(content, max_pages=self.ocr_max_pages)
            if (cache_key is not None and len(text) <= 200_000
                    and not sensitive_filename(path) and not sensitive_content(text)):
                if len(self._ocr_cache) >= 8:
                    self._ocr_cache.pop(next(iter(self._ocr_cache)))
                self._ocr_cache[cache_key] = text
            return {'text': text, 'ocr_used': True}
        if suffix in {'.txt', '.md', '.markdown', '.log', '.csv', '.json', '.xml', '.yaml', '.yml'}:
            text = content.decode('utf-8', errors='replace')
        elif suffix == '.xlsx':
            text = self._extract_xlsx(content)
        elif suffix == '.docx':
            text = self._extract_docx(content)
        elif suffix == '.odt':
            text = self._extract_odt(content)
        elif suffix == '.rtf':
            text = self._extract_rtf(content)
        else:
            text = content.decode('utf-8', errors='replace')
        return {'text': text, 'ocr_used': False}

    def get_file_text(self, path: str) -> str:
        """Backward-compatible text-only API for older callers."""
        return self.get_file_text_with_metadata(path)['text']

    def download_file(self, path: str) -> bytes:
        url = self._dav_url(path)

        response = self.session.get(
            url,
            timeout=self.timeout,
            verify=self.verify_ssl,
        )
        response.raise_for_status()

        return response.content

    def _propfind(self, path: str, depth: str = "1") -> List[Dict]:
        url = self._dav_url(path)

        body = """<?xml version="1.0" encoding="utf-8"?>
            <d:propfind xmlns:d="DAV:" xmlns:oc="http://owncloud.org/ns" xmlns:nc="http://nextcloud.org/ns">
              <d:prop>
                <d:resourcetype/>
                <d:orderby/>
                <d:getcontenttype/>
                <d:getcontentlength/>
                <d:getetag/>
                <oc:fileid/>
              </d:prop>
            </d:propfind>
            """

        response = self.session.request(
            "PROPFIND",
            url,
            data=body.encode("utf-8"),
            headers={"Depth": depth, "Content-Type": "application/xml"},
            timeout=self.timeout,
            verify=self.verify_ssl,
        )

        if response.status_code not in (207,):
            response.raise_for_status()

        root = ET.fromstring(response.content)
        items = []

        for response_el in root.findall(f"{{{self.DAV_NS}}}response"):
            href = response_el.findtext(f"{{{self.DAV_NS}}}href") or ""

            prop = None
            for propstat in response_el.findall(f"{{{self.DAV_NS}}}propstat"):
                status = propstat.findtext(f"{{{self.DAV_NS}}}status") or ""
                if " 200 " in status:
                    prop = propstat.find(f"{{{self.DAV_NS}}}prop")
                    break

            if prop is None:
                continue

            resource_type = prop.find(f"{{{self.DAV_NS}}}resourcetype")
            is_directory = (
                    resource_type is not None
                    and resource_type.find(f"{{{self.DAV_NS}}}collection") is not None
            )

            path_value = self._href_to_user_path(href)

            content_type = prop.findtext(f"{{{self.DAV_NS}}}getcontenttype") or ""
            content_length = prop.findtext(f"{{{self.DAV_NS}}}getcontentlength") or "0"
            etag = prop.findtext(f"{{{self.DAV_NS}}}getetag") or ""

            file_id = ""

            for propstat in response_el.findall(
                    f"{{{self.DAV_NS}}}propstat"
            ):
                status = propstat.findtext(
                    f"{{{self.DAV_NS}}}status"
                ) or ""

                if " 200 " not in status:
                    continue

                properties = propstat.find(
                    f"{{{self.DAV_NS}}}prop"
                )

                if properties is None:
                    continue

                file_id = (
                        properties.findtext(
                            "{http://owncloud.org/ns}fileid"
                        )
                        or properties.findtext(
                    "{http://nextcloud.org/ns}fileid"
                )
                        or ""
                ).strip()

                if file_id:
                    break

            try:
                size = int(content_length)
            except ValueError:
                size = 0

            items.append(
                {
                    "file_id": str(file_id),
                    "path": path_value,
                    "name": PurePosixPath(path_value).name,
                    "etag": etag.strip('"'),
                    "mime_type": content_type
                                 or mimetypes.guess_type(path_value)[0]
                                 or "",
                    "size": size,
                    "is_directory": is_directory,
                    "type": "directory" if is_directory else "file",
                }
            )

        return items

    def _dav_url(self, path: str) -> str:
        path = self._normalize_path(path)
        encoded_username = quote(self.username, safe="")
        encoded_path = "/".join(quote(part, safe="") for part in path.strip("/").split("/"))

        base = f"{self.base_url}/remote.php/dav/files/{encoded_username}"

        if not encoded_path:
            return base + "/"

        return f"{base}/{encoded_path}"

    def _href_to_user_path(self, href: str) -> str:
        from urllib.parse import unquote, urlparse

        parsed = urlparse(href)
        raw_path = unquote(parsed.path)

        marker = f"/remote.php/dav/files/{self.username}"
        index = raw_path.find(marker)

        if index >= 0:
            user_path = raw_path[index + len(marker):]
        else:
            user_path = raw_path

        return self._normalize_path(user_path)

    @staticmethod
    def _extract_pdf(content: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError(
                "PDF extraction requires pypdf. Add pypdf to requirements.txt."
            ) from exc

        reader = PdfReader(io.BytesIO(content))
        parts = []

        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)

        return "\n\n".join(parts)

    @staticmethod
    def _extract_xlsx(content: bytes) -> str:
        """Read cell values, not ZIP bytes; bound output for local-LLM usage."""
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines = []
        length = 0
        try:
            for sheet in workbook.worksheets[:8]:
                heading = f'Worksheet: {sheet.title}'
                lines.append(heading)
                length += len(heading)
                for row in sheet.iter_rows(max_row=100, max_col=32, values_only=True):
                    values = [str(value).strip()[:300] for value in row if value is not None]
                    if not values:
                        continue
                    line = ' | '.join(values)
                    lines.append(line)
                    length += len(line)
                    if length >= 50000:
                        return '\n'.join(lines)[:50000]
        finally:
            workbook.close()
        return '\n'.join(lines)

    @staticmethod
    def _extract_docx(content: bytes) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "DOCX extraction requires python-docx. "
                "Add python-docx to requirements.txt."
            ) from exc

        doc = Document(io.BytesIO(content))
        parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]

        for table in doc.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    parts.append(" | ".join(values))

        return "\n".join(parts)

    @staticmethod
    def _extract_odt(content: bytes) -> str:
        try:
            from odf.opendocument import load
            from odf import teletype
            from odf.text import P
        except ImportError as exc:
            raise RuntimeError(
                "ODT extraction requires odfpy. Add odfpy to requirements.txt."
            ) from exc

        document = load(io.BytesIO(content))
        return "\n".join(
            teletype.extractText(paragraph)
            for paragraph in document.getElementsByType(P)
            if teletype.extractText(paragraph).strip()
        )

    @staticmethod
    def _extract_rtf(content: bytes) -> str:
        text = content.decode("utf-8", errors="replace")
        text = re.sub(r"\\par[d]?\b", "\n", text)
        text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
        text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
        text = text.replace("{", "").replace("}", "")
        return re.sub(r"\n{3,}", "\n\n", text).strip()

    @staticmethod
    def _normalize_extension(extension: str) -> str:
        extension = str(extension).strip().lower()
        return extension if extension.startswith(".") else "." + extension

    @staticmethod
    def _normalize_path(path: str) -> str:
        path = str(path).strip().replace("\\", "/")

        if not path:
            return "/"

        if not path.startswith("/"):
            path = "/" + path

        while "//" in path:
            path = path.replace("//", "/")

        if len(path) > 1:
            path = path.rstrip("/")

        return path

    @staticmethod
    def _load_config(config_file: str) -> Dict:
        try:
            with open(config_file, "r", encoding="utf-8") as handle:
                config = yaml.safe_load(handle) or {}
        except FileNotFoundError as exc:
            raise RuntimeError(f"Configuration file not found: {config_file}") from exc
        except yaml.YAMLError as exc:
            raise RuntimeError(f"Invalid YAML configuration: {config_file}") from exc

        if not isinstance(config, dict):
            raise RuntimeError("Configuration root must be a YAML mapping.")

        return config

    def find_file_by_id(self, file_id):
        """
        Find a file anywhere in the connected user's active
        Nextcloud files using its Nextcloud file ID.

        Returns:
            dict: Current file metadata if found.
            None: Successful search with no matching file.

        Raises:
            Exception: Nextcloud or search failure.
        """

        from xml.sax.saxutils import escape

        file_id = str(file_id).strip()

        if not file_id.isdecimal():
            raise ValueError("Invalid Nextcloud file ID")

        username = escape(quote(self.username, safe=""))

        body = f"""<?xml version="1.0" encoding="UTF-8"?>
        <d:searchrequest
            xmlns:d="DAV:"
            xmlns:oc="http://owncloud.org/ns">

            <d:basicsearch>

                <d:select>
                    <d:prop>
                        <oc:fileid/>
                        <d:getcontenttype/>
                        <d:getcontentlength/>
                        <d:getetag/>
                        <d:resourcetype/>
                    </d:prop>
                </d:select>

                <d:from>
                    <d:scope>
                        <d:href>/files/{username}</d:href>
                        <d:depth>infinity</d:depth>
                    </d:scope>
                </d:from>

                <d:where>
                    <d:eq>
                        <d:prop>
                            <oc:fileid/>
                        </d:prop>
                        <d:literal>{file_id}</d:literal>
                    </d:eq>
                </d:where>
                <d:orderby/>
            </d:basicsearch>
        </d:searchrequest>
        """

        url = (
            f"{self.base_url.rstrip('/')}"
            "/remote.php/dav/"
        )

        response = self.session.request(
            "SEARCH",
            url,
            data=body.encode("utf-8"),
            headers={"Content-Type": "text/xml"},
            timeout=self.timeout,
            verify=self.verify_ssl,
        )

        if response.status_code >= 400:
            self.log.error(
                "Nextcloud SEARCH failed for file ID %s "
                "(HTTP %s): %s",
                file_id,
                response.status_code,
                response.text[:1000],
            )

            response.raise_for_status()

        if response.status_code != 207:
            raise RuntimeError(
                "Unexpected Nextcloud SEARCH response: "
                f"{response.status_code}"
            )

        root = ET.fromstring(response.content)

        dav = "{DAV:}"
        oc = "{http://owncloud.org/ns}"

        for result in root.findall(f"{dav}response"):

            href = result.findtext(f"{dav}href")

            prop = None

            for propstat in result.findall(f"{dav}propstat"):

                status = (
                        propstat.findtext(f"{dav}status") or ""
                )

                if " 200 " in status:
                    prop = propstat.find(f"{dav}prop")
                    break

            if prop is None or not href:
                raise RuntimeError(
                    "Incomplete Nextcloud SEARCH response"
                )

            returned_id = prop.findtext(
                f"{oc}fileid"
            )

            if not returned_id:
                raise RuntimeError(
                    "SEARCH response missing file ID"
                )

            if str(returned_id) != file_id:
                raise RuntimeError(
                    "SEARCH returned an unexpected file ID"
                )

            path = self._href_to_user_path(href)

            if not path or not path.startswith("/"):
                raise RuntimeError(
                    "SEARCH returned an invalid file path"
                )

            resource_type = prop.find(
                f"{dav}resourcetype"
            )

            is_directory = (
                    resource_type is not None
                    and resource_type.find(
                f"{dav}collection"
            ) is not None
            )

            return {
                "file_id": file_id,
                "path": path,
                "name": PurePosixPath(path).name,
                "etag": (
                        prop.findtext(f"{dav}getetag") or ""
                ).strip('"'),
                "mime_type": (
                        prop.findtext(
                            f"{dav}getcontenttype"
                        ) or ""
                ),
                "size": int(
                    prop.findtext(
                        f"{dav}getcontentlength"
                    ) or 0
                ),
                "is_directory": is_directory,
            }

        # Search completed successfully, but no file exists.
        return None
